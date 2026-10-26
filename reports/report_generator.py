import os
from datetime import datetime
from database.db_manager import DatabaseManager
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


class ReportGenerator:
    def __init__(self):
        self.db = DatabaseManager()
    
    def export_to_excel(self, filename="schedule_report.xlsx"):
        """Экспорт расписания в Excel"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Расписание"
        
        # Заголовок
        ws['A1'] = "Расписание занятий"
        ws['A1'].font = Font(size=16, bold=True)
        ws.merge_cells('A1:F1')
        
        ws['A2'] = f"Дата формирования: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
        ws['A2'].font = Font(size=10, italic=True)
        ws.merge_cells('A2:F2')
        
        # Заголовки столбцов
        headers = ["День", "Время", "Предмет", "Преподаватель", "Аудитория", "Группа"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Данные
        rows = self.db.execute_query("""
            SELECT s.day_of_week, s.time_slot, sub.name, u.full_name, c.number, s.group_name
            FROM schedule s
            JOIN subjects sub ON s.subject_id = sub.id
            JOIN users u ON s.teacher_id = u.id
            JOIN classrooms c ON s.classroom_id = c.id
            ORDER BY s.day_of_week, s.time_slot
        """)
        
        for row_idx, row_data in enumerate(rows, 5):
            ws.cell(row=row_idx, column=1, value=row_data['day_of_week'])
            ws.cell(row=row_idx, column=2, value=row_data['time_slot'])
            ws.cell(row=row_idx, column=3, value=row_data['name'])
            ws.cell(row=row_idx, column=4, value=row_data['full_name'])
            ws.cell(row=row_idx, column=5, value=row_data['number'])
            ws.cell(row=row_idx, column=6, value=row_data['group_name'])
        
        # Ширина столбцов
        ws.column_dimensions['A'].width = 12
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 25
        ws.column_dimensions['D'].width = 30
        ws.column_dimensions['E'].width = 15
        ws.column_dimensions['F'].width = 12
        
        # Сохраняем файл
        filepath = os.path.join("reports", filename)
        wb.save(filepath)
        
        return filepath
    
    def export_to_pdf(self, filename="schedule_report.pdf"):
        """Экспорт расписания в PDF"""
        filepath = os.path.join("reports", filename)
        
        doc = SimpleDocTemplate(filepath, pagesize=landscape(A4))
        elements = []
        
        # Заголовок
        styles = getSampleStyleSheet()
        title = Paragraph("Расписание занятий", styles['Title'])
        elements.append(title)
        
        date_text = Paragraph(f"Дата формирования: {datetime.now().strftime('%d.%m.%Y %H:%M')}", styles['Normal'])
        elements.append(date_text)
        elements.append(Paragraph("<br/><br/>", styles['Normal']))
        
        # Данные
        rows = self.db.execute_query("""
            SELECT s.day_of_week, s.time_slot, sub.name, u.full_name, c.number, s.group_name
            FROM schedule s
            JOIN subjects sub ON s.subject_id = sub.id
            JOIN users u ON s.teacher_id = u.id
            JOIN classrooms c ON s.classroom_id = c.id
            ORDER BY s.day_of_week, s.time_slot
        """)
        
        # Создаём таблицу
        table_data = [["День", "Время", "Предмет", "Преподаватель", "Аудитория", "Группа"]]
        
        for row_data in rows:
            table_data.append([
                row_data['day_of_week'],
                row_data['time_slot'],
                row_data['name'],
                row_data['full_name'],
                row_data['number'],
                row_data['group_name']
            ])
        
        table = Table(table_data, colWidths=[80, 100, 150, 180, 100, 80])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2C3E50')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 10),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.lightgrey]),
        ]))
        
        elements.append(table)
        doc.build(elements)
        
        return filepath
    
    def backup_database(self):
        """Резервное копирование базы данных"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_filename = f"school_backup_{timestamp}.db"
        backup_path = os.path.join("reports", backup_filename)
        
        # Копируем файл базы данных
        import shutil
        db_path = self.db.db_path
        shutil.copy2(db_path, backup_path)
        
        return backup_path
    
    def close(self):
        self.db.close()